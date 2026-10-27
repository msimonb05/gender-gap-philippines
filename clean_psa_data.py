from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

CSV_DIR = Path("csv")
XLSX_DIR = Path("xlsx")
OUT_DIR = Path("data")

WAGE_REGION_CSV = CSV_DIR / "Table 4. Gender Wage Gap and Average Monthly Wage Rates of Time-Rated Workers  on Full-Time Basis by Sex and Region, Philippines- August 2024.csv"
RATES_CSV = CSV_DIR / "Rates Key Employment Indicators- April 2005 to July 2026.csv"
LEVELS_CSV = CSV_DIR / "Levels of Key Employment Indicators- April 2005 to July 2026.csv"
PERSONS_CSV = CSV_DIR / "Persons in the Labor Force by Sex and by Age Group- April 2005 to July 2026.csv"
ENROLLMENT_CSV = CSV_DIR / "TABLE 10.1b Ratio of Girls to Boys in Primary, Secondary and Tertiary Education- SY 2008-2009 to SY 2023-2024.csv"
POVERTY_CSV = CSV_DIR / "Table 1. Annual Per Capita Poverty Threshold and Poverty Incidence Among Families with Measures of Precision, by Region and Province- 2018, 2021, and 2023.csv"
POVERTY_YEAR = 2023
FLEMMS_XLSX = XLSX_DIR / "2024 FLEMMS Statistical Tables Press Release Final Results.xlsx"
FLEMMS_SHEET = "Table 3"

PSA_NULLS = [".", "..", "..."]
EXPORT_DECIMALS = 1

REGION_MAP = {
    "Philippines": "PHILIPPINES",
    "National Capital Region (NCR)": "NCR",
    "National Capital Region": "NCR",
    "Cordillera Administrative Region (CAR)": "CAR",
    "Cordillera Administrative Region": "CAR",
    "Region I (Ilocos Region)": "Region I - Ilocos Region",
    "Region II (Cagayan Valley)": "Region II - Cagayan Valley",
    "Region III (Central Luzon)": "Region III - Central Luzon",
    "Region IV-A (CALABARZON)": "Region IV-A - CALABARZON",
    "MIMAROPA Region": "MIMAROPA",
    "Region V (Bicol Region)": "Region V - Bicol Region",
    "Region VI (Western Visayas)": "Region VI - Western Visayas",
    "Negros Island Region (NIR)": "NIR",
    "Region VII (Central Visayas)": "Region VII - Central Visayas",
    "Region VIII (Eastern Visayas)": "Region VIII - Eastern Visayas",
    "Region IX (Zamboanga Peninsula)": "Region IX - Zamboanga Peninsula",
    "Region X (Northern Mindanao)": "Region X - Northern Mindanao",
    "Region XI (Davao Region)": "Region XI - Davao Region",
    "Region XII (SOCCSKSARGEN)": "Region XII - SOCCSKSARGEN",
    "Region XIII (Caraga)": "Region XIII - Caraga",
    "Bangsamoro Autonomous Region in Muslim Mindanao": "BARMM",
    "Bangsamoro Autonomous Region in Muslim Mindanao (BARMM)": "BARMM",
}


def read_psa_csv(path, id_cols):
    df = pd.read_csv(path, skiprows=1, dtype=str, keep_default_na=False)
    df = df.replace(PSA_NULLS, pd.NA)
    value_cols = [c for c in df.columns if c not in id_cols]
    df[value_cols] = df[value_cols].apply(pd.to_numeric, errors="raise")
    return df


def normalize_region(series):
    cleaned = series.str.split().str.join(" ")
    mapped = cleaned.replace(REGION_MAP)
    return mapped


def to_survey_months(df, required_col):
    df = df[df["Month"] != "Annual"].copy()
    df = df.dropna(subset=[required_col])
    df["date"] = pd.to_datetime(df["Year"] + " " + df["Month"], format="%Y %B")
    assert df["date"].is_unique, "duplicate survey dates"
    assert df["date"].is_monotonic_increasing, "survey dates out of order"
    return df


def clean_wage_gap():
    df = pd.read_csv(WAGE_REGION_CSV, skiprows=1)
    df.columns = ["region", "wage_male", "wage_female", "gap_pct"]

    m, f = df["wage_male"], df["wage_female"]
    recomputed = (m - f) / m * 100
    gap_rounding = 0.05
    wage_rounding = 100 * 0.5 * (1 / m + f / m**2)
    err = (recomputed - df["gap_pct"]).abs()
    assert (err <= gap_rounding + wage_rounding).all(), (
        f"gap_pct is not (male - female) / male * 100 beyond rounding: max error {err.max()}"
    )

    national = df[df["region"] == "PHILIPPINES"]
    assert len(national) == 1, "expected exactly one PHILIPPINES row"
    df = df[df["region"] != "PHILIPPINES"].copy()
    df["region"] = normalize_region(df["region"])
    assert len(df) == 17, f"expected 17 regions, got {len(df)}"
    assert df["region"].is_unique, "duplicate regions after normalization"

    df.to_csv(OUT_DIR / "wage_gap_by_region.csv", index=False)
    pd.DataFrame({"national_gap_pct": national["gap_pct"].to_list()}).to_csv(
        OUT_DIR / "national_wage_gap.csv", index=False
    )
    return df


def assert_lfpr_matches_levels(rates, levels):
    merged = rates.merge(levels, on=["Year", "Month"], how="left", validate="one_to_one")
    for sex in ["Male", "Female"]:
        recomputed = (
            merged[f"Persons in the Labor Force {sex}"]
            / merged[f"Total Population 15 Years Old and Over {sex}"]
            * 100
        )
        max_err = (recomputed - merged[f"Labor Force Participation Rate {sex}"]).abs().max()
        assert max_err <= 0.0015, f"LFPR {sex} does not match Levels file: max error {max_err}"
    return merged


def clean_labor_participation(levels):
    rates = read_psa_csv(RATES_CSV, ["Year", "Month"])
    rates = to_survey_months(rates, "Labor Force Participation Rate Female")
    assert_lfpr_matches_levels(rates, levels)

    out = rates[["date", "Labor Force Participation Rate Male", "Labor Force Participation Rate Female"]].copy()
    out.columns = ["date", "lfpr_male", "lfpr_female"]
    assert out[["lfpr_male", "lfpr_female"]].notna().all().all(), "missing LFPR after dropping empty months"
    assert out[["lfpr_male", "lfpr_female"]].gt(0).all().all(), "LFPR <= 0"
    assert out[["lfpr_male", "lfpr_female"]].lt(100).all().all(), "LFPR >= 100"

    spacing = (out["date"].dt.year * 12 + out["date"].dt.month).diff()
    print(f"  survey spacing (months): min {spacing.min():.0f}, max {spacing.max():.0f}")
    print(f"  first 1-month step: {out.loc[spacing.eq(1), 'date'].min():%Y-%m}")

    out["date"] = out["date"].dt.strftime("%Y-%m")
    out.to_csv(OUT_DIR / "labor_participation.csv", index=False)
    return out


def clean_labor_force_persons(levels):
    persons = read_psa_csv(PERSONS_CSV, ["Year", "Month"])
    persons = to_survey_months(persons, "Female Total")

    merged = persons.merge(levels, on=["Year", "Month"], how="left", validate="one_to_one")
    for sex in ["Male", "Female"]:
        diff = (merged[f"{sex} Total"] - merged[f"Persons in the Labor Force {sex}"]).abs().max()
        assert diff == 0, f"{sex} Total differs from Levels file by up to {diff} thousand"

    out = persons[["date", "Male Total", "Female Total"]].copy()
    out.columns = ["date", "male_thousands", "female_thousands"]
    assert out[["male_thousands", "female_thousands"]].notna().all().all(), "missing labor force counts"
    out["date"] = out["date"].dt.strftime("%Y-%m")
    out.to_csv(OUT_DIR / "labor_force_persons.csv", index=False)
    return out


def clean_employment_annual(levels):
    rates = read_psa_csv(RATES_CSV, ["Year", "Month"])
    rates = rates[rates["Month"] == "Annual"].dropna(subset=["Labor Force Participation Rate Female"])
    merged = assert_lfpr_matches_levels(rates, levels)

    columns = {"Year": "year"}
    for sex, s in [("Male", "male"), ("Female", "female")]:
        columns |= {
            f"Total Population 15 Years Old and Over {sex}": f"pop15_{s}_thousands",
            f"Persons in the Labor Force {sex}": f"lf_{s}_thousands",
            f"Labor Force Participation Rate {sex}": f"lfpr_{s}",
            f"Employment Rate {sex}": f"emp_rate_{s}",
            f"Unemployment Rate {sex}": f"unemp_rate_{s}",
            f"Underemployment Rate {sex}": f"underemp_rate_{s}",
        }
    out = merged[list(columns)].rename(columns=columns)
    out["year"] = out["year"].astype(int)

    assert out["year"].to_list() == list(range(2005, 2026)), f"unexpected annual years: {out['year'].to_list()}"
    assert out.notna().all().all(), "missing values in annual employment table"
    for s in ["male", "female"]:
        total = out[f"emp_rate_{s}"] + out[f"unemp_rate_{s}"]
        assert (total - 100).abs().max() <= 0.0015, f"employment + unemployment rate != 100 for {s}"

    out.to_csv(OUT_DIR / "employment_annual.csv", index=False)
    return out


def clean_poverty():
    df = read_psa_csv(POVERTY_CSV, ["Geolocation"])
    y = POVERTY_YEAR
    columns = {
        "Geolocation": "region",
        f"Poverty Incidence Among Families (%) {y}": "poverty_incidence_pct",
        f"Standard Error {y}": "se_pct",
        f"95% Confidence Interval (Lower Limits) {y}": "ci_low_pct",
        f"95% Confidence Interval (Upper Limits) {y}": "ci_high_pct",
    }
    df = df[list(columns)].rename(columns=columns)

    national = df[df["region"] == "PHILIPPINES"]
    assert len(national) == 1, "expected exactly one PHILIPPINES row"
    regions = df[df["region"].str.match(r"^\.\.[^.]")].copy()
    regions["region"] = normalize_region(regions["region"].str.removeprefix(".."))

    assert len(regions) == 18, f"expected 18 regions, got {len(regions)}"
    assert regions["region"].is_unique, "duplicate regions after normalization"
    assert regions.notna().all().all(), "missing poverty values for a region"
    z = 1.96
    rounding = 0.5 + z * 0.05 + 0.05
    est, se = regions["poverty_incidence_pct"], regions["se_pct"]
    assert ((est - z * se) - regions["ci_low_pct"]).abs().le(rounding).all(), "CI lower limit != est - 1.96 SE"
    assert ((est + z * se) - regions["ci_high_pct"]).abs().le(rounding).all(), "CI upper limit != est + 1.96 SE"

    regions.to_csv(OUT_DIR / "poverty_by_region.csv", index=False)
    national.to_csv(OUT_DIR / "national_poverty.csv", index=False)
    return regions


def find_group_col(header_row, prefix):
    matches = [i for i, v in enumerate(header_row) if isinstance(v, str) and v.startswith(prefix)]
    assert len(matches) == 1, f"expected one header starting with {prefix!r}, found {len(matches)}"
    return matches[0]


def clean_literacy():
    wb = load_workbook(FLEMMS_XLSX, read_only=True, data_only=True)
    rows = list(wb[FLEMMS_SHEET].iter_rows(values_only=True))
    wb.close()

    group_idx = [i for i, r in enumerate(rows) if "Region and Age Group" in r]
    assert len(group_idx) == 1, "header row with 'Region and Age Group' not found exactly once"
    header_i = group_idx[0]
    group_row, sex_row = rows[header_i], rows[header_i + 1]

    label_col = group_row.index("Region and Age Group")
    pop_col = find_group_col(group_row, "Population 10 to 64")
    lit_col = find_group_col(group_row, "Number of Functionally Literate")
    rate_col = find_group_col(group_row, "Functional Literacy Rate")
    for col in [pop_col, lit_col, rate_col]:
        assert sex_row[col:col + 3] == ("Both Sexes", "Male", "Female"), f"unexpected sex columns at {col}"

    records = []
    for row in rows[header_i + 2:]:
        label = row[label_col]
        if label is None:
            continue
        if label == "Age Group":
            break
        records.append({
            "region": label,
            "pop_both": row[pop_col], "pop_male": row[pop_col + 1], "pop_female": row[pop_col + 2],
            "lit_both": row[lit_col], "lit_male": row[lit_col + 1], "lit_female": row[lit_col + 2],
            "func_lit_rate_both": row[rate_col],
            "func_lit_rate_male": row[rate_col + 1],
            "func_lit_rate_female": row[rate_col + 2],
        })

    df = pd.DataFrame(records)
    for sex in ["both", "male", "female"]:
        recomputed = df[f"lit_{sex}"] / df[f"pop_{sex}"] * 100
        max_err = (recomputed - df[f"func_lit_rate_{sex}"]).abs().max()
        assert max_err < 1e-6, f"literacy rate {sex} != literate / population * 100: max error {max_err}"
        assert df[f"func_lit_rate_{sex}"].between(0, 100).all(), f"literacy rate {sex} outside 0-100"

    df["region"] = normalize_region(df["region"])
    rate_cols = ["func_lit_rate_both", "func_lit_rate_male", "func_lit_rate_female"]
    df = df[["region"] + rate_cols]
    df[rate_cols] = df[rate_cols].round(EXPORT_DECIMALS)

    national = df[df["region"] == "PHILIPPINES"]
    assert len(national) == 1, "expected exactly one Philippines row"
    df = df[df["region"] != "PHILIPPINES"].copy()
    assert len(df) == 18, f"expected 18 regions, got {len(df)}"
    assert df["region"].is_unique, "duplicate regions after normalization"

    df.to_csv(OUT_DIR / "literacy_rate.csv", index=False)
    national.to_csv(OUT_DIR / "national_literacy.csv", index=False)
    return df


def clean_school_enrollment():
    df = read_psa_csv(ENROLLMENT_CSV, ["Level"])
    df_long = df.melt(id_vars=["Level"], var_name="school_year", value_name="ratio")
    df_long = df_long.dropna(subset=["ratio"])
    df_long["year_start"] = df_long["school_year"].str.split("-").str[0].astype(int)

    counts = df_long["Level"].value_counts().to_dict()
    assert counts == {"Primary": 16, "Secondary": 16, "Tertiary": 6}, f"unexpected rows per level: {counts}"
    assert df_long["ratio"].gt(0).all(), "non-positive girls-to-boys ratio"

    df_long.to_csv(OUT_DIR / "school_enrollment.csv", index=False)
    return df_long


def main():
    OUT_DIR.mkdir(exist_ok=True)
    levels = read_psa_csv(LEVELS_CSV, ["Year", "Month"])

    print("wage gap")
    wage = clean_wage_gap()
    print("labor participation")
    lfpr = clean_labor_participation(levels)
    print("labor force persons")
    persons = clean_labor_force_persons(levels)
    print("literacy")
    literacy = clean_literacy()
    print("school enrollment")
    enrollment = clean_school_enrollment()
    print("employment annual")
    clean_employment_annual(levels)
    print("poverty")
    poverty = clean_poverty()

    assert lfpr["date"].equals(persons["date"]), "LFPR and labor force files cover different dates"
    missing_in_wage = set(literacy["region"]) - set(wage["region"])
    assert missing_in_wage == {"NIR"}, f"region names disagree between files: {missing_in_wage}"
    assert set(wage["region"]) <= set(literacy["region"]), "wage regions missing from literacy"
    assert set(poverty["region"]) == set(literacy["region"]), "poverty and literacy regions disagree"

    for name in sorted(p.name for p in OUT_DIR.glob("*.csv")):
        shape = pd.read_csv(OUT_DIR / name).shape
        print(f"{name:30s} {shape[0]:4d} rows x {shape[1]} cols")


if __name__ == "__main__":
    main()
